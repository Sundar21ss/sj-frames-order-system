import os
import sqlite3
import secrets
from datetime import datetime
from pathlib import Path
from functools import wraps
from io import BytesIO

from flask import (
    Flask, render_template, request, redirect, url_for, session, flash,
    send_from_directory, abort, make_response, jsonify
)
from werkzeug.utils import secure_filename
from PIL import Image, UnidentifiedImageError

BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = Path(os.environ.get('DATA_DIR', str(BASE_DIR))).resolve()
DATA_DIR.mkdir(parents=True, exist_ok=True)
UPLOAD_DIR = DATA_DIR / 'private_uploads'
UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
DB_PATH = DATA_DIR / 'orders.db'
ALLOWED_EXTENSIONS = {'jpg', 'jpeg', 'png', 'webp'}
MAX_FILE_BYTES = 10 * 1024 * 1024
ALLOWED_STATUSES = {
    'New', 'Reviewing', 'Awaiting Customer Approval',
    'In Production', 'Completed', 'Cancelled'
}

app = Flask(__name__)
app.config.update(
    SECRET_KEY=os.environ.get('SECRET_KEY') or secrets.token_hex(32),
    MAX_CONTENT_LENGTH=22 * 1024 * 1024,
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE='Lax',
    SESSION_COOKIE_SECURE=os.environ.get('COOKIE_SECURE', '0') == '1',
    PERMANENT_SESSION_LIFETIME=3600,
)
ADMIN_USERNAME = os.environ.get('ADMIN_USERNAME', 'admin')
# Never ship a usable hard-coded admin password. Set ADMIN_PASSWORD in the environment.
ADMIN_PASSWORD = os.environ.get('ADMIN_PASSWORD', '')
# WhatsApp number must include the country calling code for wa.me links.
# Default: India (+91) + business mobile number 9150900482.
WHATSAPP_BUSINESS_NUMBER = ''.join(
    ch for ch in os.environ.get('WHATSAPP_BUSINESS_NUMBER', '919150900482') if ch.isdigit()
)

# Refuse unsafe default credentials in a production deployment.
if os.environ.get('APP_ENV', '').lower() == 'production':
    if not os.environ.get('SECRET_KEY') or len(os.environ.get('SECRET_KEY', '1234567890987654321')) < 32:
        raise RuntimeError('Set SECRET_KEY to a persistent random value of at least 32 characters.')
    if not os.environ.get('ADMIN_PASSWORD') or len(ADMIN_PASSWORD) < 14:
        raise RuntimeError('Set ADMIN_PASSWORD to a unique password of at least 14 characters.')


def db():
    conn = sqlite3.connect(DB_PATH, timeout=15)
    conn.row_factory = sqlite3.Row
    return conn


def init_db():
    with db() as conn:
        conn.execute('''CREATE TABLE IF NOT EXISTS orders (
            id INTEGER PRIMARY KEY AUTOINCREMENT, reference TEXT UNIQUE NOT NULL,
            created_at TEXT NOT NULL, customer_name TEXT NOT NULL, whatsapp TEXT NOT NULL,
            brand TEXT NOT NULL, model TEXT NOT NULL, lid_length TEXT NOT NULL, lid_width TEXT NOT NULL,
            lid_photo TEXT NOT NULL, features TEXT, design_image TEXT NOT NULL, finish TEXT NOT NULL,
            instructions TEXT, status TEXT NOT NULL DEFAULT 'New')''')


def csrf_token():
    token = session.get('_csrf_token')
    if not token:
        token = secrets.token_urlsafe(32)
        session['_csrf_token'] = token
    return token


app.jinja_env.globals['csrf_token'] = csrf_token


@app.before_request
def protect_post_requests():
    if request.method == 'POST':
        submitted = request.form.get('_csrf_token', '')
        expected = session.get('_csrf_token', '')
        if not submitted or not expected or not secrets.compare_digest(submitted, expected):
            abort(400, description='Your session expired or the form was invalid. Please reload and try again.')


@app.after_request
def security_headers(response):
    response.headers.setdefault('X-Content-Type-Options', 'nosniff')
    response.headers.setdefault('X-Frame-Options', 'DENY')
    response.headers.setdefault('Referrer-Policy', 'strict-origin-when-cross-origin')
    response.headers.setdefault('Cache-Control', 'no-store')
    response.headers.setdefault('Content-Security-Policy', "default-src 'self' 'unsafe-inline' data:; img-src 'self' data:; form-action 'self'; frame-ancestors 'none'; base-uri 'self'")
    return response


def allowed(file):
    return bool(file and file.filename and '.' in file.filename and file.filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS)


def save_image(file, prefix):
    if not allowed(file):
        raise ValueError('Please upload JPG, PNG or WebP images only.')
    file.stream.seek(0, os.SEEK_END)
    size = file.stream.tell()
    file.stream.seek(0)
    if size == 0:
        raise ValueError('The selected image is empty. Please choose another image.')
    if size > MAX_FILE_BYTES:
        raise ValueError('Each image must be 10 MB or smaller.')

    ext = file.filename.rsplit('.', 1)[1].lower()
    try:
        with Image.open(file.stream) as img:
            fmt = (img.format or '').upper()
            expected_formats = {'jpg': {'JPEG'}, 'jpeg': {'JPEG'}, 'png': {'PNG'}, 'webp': {'WEBP'}}
            if fmt not in expected_formats[ext]:
                raise ValueError('The file contents do not match the image extension.')
            img.verify()
        file.stream.seek(0)
        with Image.open(file.stream) as img:
            img.load()
            if img.width > 8000 or img.height > 8000:
                raise ValueError('Images must be no larger than 8000 × 8000 pixels.')
            # Re-encode the image to remove unexpected embedded payloads and metadata.
            if ext in {'jpg', 'jpeg'}:
                if img.mode not in ('RGB', 'L'):
                    img = img.convert('RGB')
                output_format, suffix = 'JPEG', 'jpg'
            elif ext == 'png':
                output_format, suffix = 'PNG', 'png'
            else:
                if img.mode not in ('RGB', 'RGBA'):
                    img = img.convert('RGBA' if 'transparency' in img.info else 'RGB')
                output_format, suffix = 'WEBP', 'webp'
            name = secure_filename(f'{prefix}_{secrets.token_hex(16)}.{suffix}')
            target = UPLOAD_DIR / name
            img.save(target, format=output_format)
            return name
    except (UnidentifiedImageError, OSError, Image.DecompressionBombError) as exc:
        raise ValueError('That file is not a valid supported image. Please upload a JPG, PNG or WebP image.') from exc


def remove_saved_images(*names):
    for name in names:
        if name:
            try:
                (UPLOAD_DIR / name).unlink(missing_ok=True)
            except OSError:
                app.logger.exception('Could not remove an unused uploaded image')


@app.get('/')
def index():
    success_reference = session.pop('success_reference', None)
    success_customer_name = session.pop('success_customer_name', '')
    success_model = session.pop('success_model', '')
    success_customer_whatsapp = session.pop('success_customer_whatsapp', '')
    whatsapp_message = ''
    if success_reference:
        whatsapp_message = (
            f'Hello SJ FRAMES! I have submitted my laptop skin order.\n'
            f'Order reference: {success_reference}\n'
            f'Customer name: {success_customer_name}\n'
            f'Laptop model: {success_model}\n'
            f'Customer WhatsApp: {success_customer_whatsapp}\n'
            'Please confirm that you received my order. Thank you.'
        )
    return render_template('index.html', success_reference=success_reference,
                           success_customer_name=success_customer_name, success_model=success_model,
                           success_customer_whatsapp=success_customer_whatsapp,
                           whatsapp_message=whatsapp_message, whatsapp_business_number=WHATSAPP_BUSINESS_NUMBER,
                           form_values={}, field_errors={})


def render_order_form_error(message, field_errors=None, status=200):
    # Preserve typed text/select/checkbox values after server-side validation.
    # Browsers intentionally do not allow restoring selected file inputs.
    form_values = request.form.to_dict(flat=True)
    errors = field_errors or {}
    # AJAX validation keeps the original page and file inputs alive in the browser.
    if request.headers.get('Accept', '').find('application/json') >= 0:
        return jsonify(ok=False, message=message, field_errors=errors), status
    return render_template(
        'index.html', success_reference=None, success_customer_name='',
        success_model='', success_customer_whatsapp='', whatsapp_message='',
        whatsapp_business_number=WHATSAPP_BUSINESS_NUMBER,
        error_message=message, field_errors=errors, form_values=form_values,
    ), status


@app.post('/submit')
def submit():
    form = request.form
    required = ['customerName', 'whatsapp', 'brand', 'model', 'length', 'width', 'finish',
                'detailsConfirm', 'imageRights', 'previewApproval', 'privacyConsent']
    if any(not form.get(k, '').strip() for k in required) or any(
        form.get(k) != 'yes' for k in ['detailsConfirm', 'imageRights', 'previewApproval', 'privacyConsent']
    ) or 'lidPhoto' not in request.files or 'designImage' not in request.files:
        message = 'Please complete the required fields, accept each confirmation, and upload both images.'
        return render_order_form_error(message, {'general': message})

    try:
        length = float(form['length'])
        width = float(form['width'])
        if not (5 <= length <= 100 and 5 <= width <= 100):
            raise ValueError('Laptop lid measurements must be between 5 and 100 cm.')
    except (TypeError, ValueError) as exc:
        message = str(exc) or 'Please enter valid laptop lid measurements.'
        return render_order_form_error(message, {'measurements': message})

    customer_name = form['customerName'].strip()
    whatsapp = form['whatsapp'].strip()
    model = form['model'].strip()
    if not customer_name or len(customer_name) > 120 or len(whatsapp) > 30 or not model or len(model) > 120:
        message = 'Please check the customer name, WhatsApp number and laptop model.'
        return render_order_form_error(message, {'customerName': message, 'whatsapp': message, 'model': message})
    digits = ''.join(ch for ch in whatsapp if ch.isdigit())
    if not (8 <= len(digits) <= 15) or not all(ch in '+0123456789() -' for ch in whatsapp):
        message = 'Please enter a valid WhatsApp number (8–15 digits; use only digits, spaces, +, brackets or hyphens).'
        return render_order_form_error(message, {'whatsapp': message})

    lid_name = design_name = None
    try:
        lid_name = save_image(request.files['lidPhoto'], 'lid')
        design_name = save_image(request.files['designImage'], 'design')
        reference = 'SJ-' + datetime.now().strftime('%y%m%d') + '-' + secrets.token_hex(3).upper()
        with db() as conn:
            conn.execute('''INSERT INTO orders (reference,created_at,customer_name,whatsapp,brand,model,lid_length,lid_width,lid_photo,features,design_image,finish,instructions)
                VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)''', (
                reference, datetime.now().strftime('%Y-%m-%d %H:%M:%S'), customer_name, whatsapp,
                form['brand'].strip()[:80], model, str(length), str(width), lid_name,
                form.get('features', '').strip()[:2000], design_name,
                form['finish'].strip()[:40], form.get('instructions', '').strip()[:4000]))
        try:
            sync_excel_database()
        except Exception:
            app.logger.exception('Order saved but Excel mirror could not be refreshed')
        session['success_reference'] = reference
        session['success_customer_name'] = customer_name
        session['success_model'] = model
        session['success_customer_whatsapp'] = whatsapp
        if request.headers.get('Accept', '').find('application/json') >= 0:
            return jsonify(ok=True, reference=reference)
        return redirect(url_for('index'))
    except ValueError as exc:
        remove_saved_images(lid_name, design_name)
        return render_order_form_error(str(exc), {'general': str(exc)})
    except sqlite3.Error:
        remove_saved_images(lid_name, design_name)
        app.logger.exception('Order database insert failed')
        message = 'We could not save your order right now. Please try again later or contact SJ Frames on WhatsApp.'
        return render_order_form_error(message, {'general': message}, 500)
    except OSError:
        remove_saved_images(lid_name, design_name)
        app.logger.exception('Order file storage failed')
        message = 'We could not save the uploaded images. Please try again later.'
        return render_order_form_error(message, {'general': message}, 500)


@app.get('/admin/login')
def admin_login():
    return render_template('login.html', error=None)


@app.post('/admin/login')
def admin_login_post():
    submitted_user = request.form.get('username', '')
    submitted_password = request.form.get('password', '')
    if ADMIN_PASSWORD and secrets.compare_digest(submitted_user, ADMIN_USERNAME) and secrets.compare_digest(submitted_password, ADMIN_PASSWORD):
        session.clear()
        session['admin'] = True
        session.permanent = True
        csrf_token()
        return redirect(url_for('admin'))
    return render_template('login.html', error='Incorrect username or password.'), 401


def admin_required(fn):
    @wraps(fn)
    def wrapper(*args, **kwargs):
        if not session.get('admin'):
            return redirect(url_for('admin_login'))
        return fn(*args, **kwargs)
    return wrapper


@app.get('/admin')
@admin_required
def admin():
    with db() as conn:
        orders = conn.execute('SELECT * FROM orders ORDER BY id DESC').fetchall()
        total = conn.execute('SELECT COUNT(*) FROM orders').fetchone()[0]
        counts = {row['status']: row['n'] for row in conn.execute('SELECT status, COUNT(*) AS n FROM orders GROUP BY status').fetchall()}
    return render_template('admin.html', orders=orders, total=total, counts=counts,
                           statuses=sorted(ALLOWED_STATUSES), whatsapp_business_number=WHATSAPP_BUSINESS_NUMBER)


@app.get('/admin/export.xlsx')
@admin_required
def export_orders_excel():
    """Export order records to a formatted Excel workbook; customer images are not embedded."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    columns = [
        ('Order ID', 'id'), ('Reference Number', 'reference'), ('Submitted At', 'created_at'),
        ('Customer Name', 'customer_name'), ('Customer WhatsApp', 'whatsapp'),
        ('Laptop Brand', 'brand'), ('Laptop Model', 'model'), ('Lid Length (cm)', 'lid_length'),
        ('Lid Width (cm)', 'lid_width'), ('Logo / Cut-out Details', 'features'),
        ('Preferred Finish', 'finish'), ('Design Instructions', 'instructions'), ('Order Status', 'status')
    ]
    with db() as conn:
        rows = conn.execute('SELECT * FROM orders ORDER BY id DESC').fetchall()

    wb = Workbook()
    ws = wb.active
    ws.title = 'Orders'
    ws.append([label for label, _ in columns])
    for order in rows:
        ws.append([order[key] or '' for _, key in columns])
    header_fill = PatternFill('solid', fgColor='6D28D9')
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = Font(bold=True, color='FFFFFF')
        cell.alignment = Alignment(vertical='center', wrap_text=True)
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    widths = [10, 23, 21, 24, 20, 18, 28, 16, 15, 36, 20, 45, 28]
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    ws.sheet_view.showGridLines = False

    summary = wb.create_sheet('Summary')
    summary.append(['SJ FRAMES · ORDER SUMMARY', 'Count'])
    summary.append(['Total orders', len(rows)])
    for status in sorted(ALLOWED_STATUSES):
        summary.append([status, sum(1 for row in rows if row['status'] == status)])
    for cell in summary[1]:
        cell.fill = header_fill
        cell.font = Font(bold=True, color='FFFFFF')
    summary.column_dimensions['A'].width = 34
    summary.column_dimensions['B'].width = 16
    summary.freeze_panes = 'A2'
    summary.sheet_view.showGridLines = False

    output = BytesIO()
    wb.save(output)
    filename = 'SJ_FRAMES_Orders_' + datetime.now().strftime('%Y%m%d_%H%M') + '.xlsx'
    response = make_response(output.getvalue())
    response.headers['Content-Type'] = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    response.headers['Content-Disposition'] = f'attachment; filename={filename}'
    response.headers['Cache-Control'] = 'private, no-store'
    return response


@app.post('/admin/status/<int:order_id>')
@admin_required
def update_status(order_id):
    status = request.form.get('status', 'New')
    if status not in ALLOWED_STATUSES:
        abort(400)
    with db() as conn:
        cursor = conn.execute('UPDATE orders SET status=? WHERE id=?', (status, order_id))
        if cursor.rowcount == 0:
            abort(404)
    try:
        sync_excel_database()
    except Exception:
        app.logger.exception('Order status updated but Excel mirror could not be refreshed')
        flash('Status updated, but the Excel mirror could not refresh. Use Export Excel to generate a fresh copy.')
    else:
        flash('Order status updated.')
    return redirect(url_for('admin'))


@app.post('/admin/logout')
@admin_required
def admin_logout():
    session.clear()
    return redirect(url_for('admin_login'))


@app.get('/admin/image/<path:filename>')
@admin_required
def admin_image(filename):
    # Upload names are generated by this app; only serve a basename from private storage.
    if Path(filename).name != filename:
        abort(404)
    response = make_response(send_from_directory(UPLOAD_DIR, filename, conditional=True))
    response.headers['Cache-Control'] = 'private, no-store'
    response.headers['X-Content-Type-Options'] = 'nosniff'
    return response


@app.errorhandler(400)
def bad_request(error):
    if request.path.startswith('/admin'):
        return 'Invalid or expired form. Please go back, reload the page and try again.', 400
    return render_template('index.html', success_reference=None, error_message='Invalid or expired form. Please reload the page and try again.', form_values={}, field_errors={}), 400


@app.errorhandler(413)
def too_large(_error):
    return render_template('index.html', success_reference=None, error_message='Upload too large. Please keep each image under 10 MB.', form_values={}, field_errors={}), 413


@app.errorhandler(500)
def internal_error(_error):
    app.logger.exception('Unexpected server error')
    return 'Something went wrong. Please try again later.', 500


def sync_excel_database():
    """Maintain a convenient Excel mirror of SQLite; SQLite remains the source of truth."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    columns = [
        ('Order ID', 'id'), ('Reference Number', 'reference'), ('Submitted At', 'created_at'),
        ('Customer Name', 'customer_name'), ('Customer WhatsApp', 'whatsapp'),
        ('Laptop Brand', 'brand'), ('Laptop Model', 'model'), ('Lid Length (cm)', 'lid_length'),
        ('Lid Width (cm)', 'lid_width'), ('Logo / Cut-out Details', 'features'),
        ('Preferred Finish', 'finish'), ('Design Instructions', 'instructions'), ('Order Status', 'status')
    ]
    with db() as conn:
        rows = conn.execute('SELECT * FROM orders ORDER BY id DESC').fetchall()
    wb = Workbook()
    ws = wb.active
    ws.title = 'Orders'
    ws.append([label for label, _ in columns])
    for order in rows:
        ws.append([order[key] or '' for _, key in columns])
    fill = PatternFill('solid', fgColor='6D28D9')
    for cell in ws[1]:
        cell.fill = fill
        cell.font = Font(bold=True, color='FFFFFF')
        cell.alignment = Alignment(vertical='center', wrap_text=True)
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    widths = [10, 23, 21, 24, 20, 18, 28, 16, 15, 36, 20, 45, 28]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    ws.sheet_view.showGridLines = False
    summary = wb.create_sheet('Summary')
    summary.append(['SJ FRAMES · ORDER SUMMARY', 'Count'])
    summary.append(['Total orders', len(rows)])
    for status in sorted(ALLOWED_STATUSES):
        summary.append([status, sum(1 for order in rows if order['status'] == status)])
    for cell in summary[1]:
        cell.fill = fill
        cell.font = Font(bold=True, color='FFFFFF')
    summary.column_dimensions['A'].width = 34
    summary.column_dimensions['B'].width = 16
    summary.freeze_panes = 'A2'
    summary.sheet_view.showGridLines = False
    target = DATA_DIR / 'orders.xlsx'
    temporary = DATA_DIR / 'orders.xlsx.tmp'
    with temporary.open('wb') as stream:
        wb.save(stream)
    temporary.replace(target)



init_db()
try:
    sync_excel_database()
except Exception:
    app.logger.exception('Could not initialize Excel mirror; the SQLite database remains available')
if __name__ == '__main__':
    app.run(debug=os.environ.get('FLASK_DEBUG') == '1')
